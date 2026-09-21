import collections
import csv
import difflib
import io
import re

from openpyxl import Workbook, load_workbook

from db import PALETTE, query, tx

KINDS = ['lessons', 'teachers', 'classes', 'subjects']


def read_text(text):
    delimiter = '\t' if '\t' in text else ','
    return list(csv.reader(io.StringIO(text), delimiter=delimiter))


def read_file(name, content):
    if name.lower().endswith('.xlsx'):
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        return [['' if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    return read_text(content.decode('utf-8-sig'))


def to_int(value, default=None):
    if value in ('', None):
        return default
    try:
        return int(float(value))
    except ValueError:
        raise ValueError(f'"{value}" is not a number')


def parse_row(kind, v):
    if kind != 'lessons':
        if not v.get('name'):
            raise ValueError('Name is empty')
        return [v]
    for key in ('class', 'teacher', 'subject'):
        if not v.get(key):
            raise ValueError(f'{key.capitalize()} is empty')
    base = {key: v[key] for key in ('class', 'teacher', 'subject')}
    if v.get('pattern'):
        counts = collections.Counter(to_int(x.strip()) for x in v['pattern'].split('+'))
        for length in counts:
            if length not in (1, 2):
                raise ValueError(f'Lesson length {length} is not 1 or 2')
        return [dict(base, length=length, count=n) for length, n in sorted(counts.items(), reverse=True)]
    length = to_int(v.get('length'), 1)
    count = to_int(v.get('count'), 1)
    if length not in (1, 2):
        raise ValueError(f'Lesson length {length} is not 1 or 2')
    if count < 1:
        raise ValueError('Lessons per week must be at least 1')
    return [dict(base, length=length, count=count)]


def parse(kind, rows, header, mapping):
    if kind == 'lessons':
        for key in ('class', 'teacher', 'subject'):
            if key not in mapping:
                raise ValueError(f'Choose a column for {key}')
        if 'pattern' not in mapping and 'count' not in mapping:
            raise ValueError('Choose a column for lessons per week or periods pattern')
    elif 'name' not in mapping:
        raise ValueError('Choose a column for name')
    items, errors = [], []
    for i in range(1 if header else 0, len(rows)):
        row = rows[i]
        if not any(str(x).strip() for x in row):
            continue
        v = {m: str(row[j]).strip() for j, m in enumerate(mapping) if m != 'ignore' and j < len(row)}
        try:
            items += parse_row(kind, v)
        except ValueError as e:
            errors.append({'row': i + 1, 'message': str(e)})
    return items, errors


def entity_keys(kind):
    if kind == 'lessons':
        return {'classes': 'class', 'teachers': 'teacher', 'subjects': 'subject'}
    return {kind: 'name'}


def normal(s):
    return re.sub(r'[^a-z0-9]', '', s.lower())


def suggest(name, existing):
    best, score = None, 0.0
    for e in existing:
        r = difflib.SequenceMatcher(None, normal(name), normal(e['name'])).ratio()
        if r > score:
            best, score = e['id'], r
    return best if score >= 0.8 else None


def unmatched(kind, items):
    out = {}
    for table, key in entity_keys(kind).items():
        existing = query(f'select id, name from {table}')
        known = {e['name'] for e in existing}
        seen = []
        for item in items:
            if item[key] not in known and item[key] not in seen:
                seen.append(item[key])
        out[table] = [{'name': n, 'suggest_id': suggest(n, existing)} for n in seen]
    return out


def unique_short(c, base):
    base = base or 'T'
    short, i = base, 1
    while c.execute('select 1 from teachers where short = ?', (short,)).fetchone():
        i += 1
        short = f'{base}{i}'
    return short


def initials(name):
    return ''.join(w[0] for w in re.split(r'[\s.]+', name) if w).upper()


def commit(kind, items, decisions):
    created, linked = collections.Counter(), collections.Counter()
    added = updated = 0
    cache = {}
    with tx() as c:

        def resolve(table, name, extra):
            key = (table, name)
            if key in cache:
                return cache[key]
            row = c.execute(f'select id from {table} where name = ?', (name,)).fetchone()
            decision = decisions.get(table, {}).get(name, {})
            if row:
                item_id = row['id']
                linked[table] += 1
            elif decision.get('action') == 'link' and decision.get('id'):
                item_id = decision['id']
                linked[table] += 1
            else:
                item_id = create(table, name, extra)
                created[table] += 1
                extra = {}
            update(table, item_id, extra)
            cache[key] = item_id
            return item_id

        def create(table, name, extra):
            if table == 'teachers':
                short = unique_short(c, extra.get('short') or initials(name))
                return c.execute('insert into teachers (name, short) values (?, ?)', (name, short)).lastrowid
            if table == 'subjects':
                count = c.execute('select count(*) from subjects').fetchone()[0]
                return c.execute(
                    'insert into subjects (name, short, color) values (?, ?, ?)',
                    (name, extra.get('short') or name[:3].upper(), PALETTE[count % len(PALETTE)]),
                ).lastrowid
            return c.execute('insert into classes (name, label) values (?, ?)', (name, extra.get('label', ''))).lastrowid

        def update(table, item_id, extra):
            if table == 'teachers' and extra.get('short'):
                taken = c.execute('select 1 from teachers where short = ? and id != ?', (extra['short'], item_id)).fetchone()
                if not taken:
                    c.execute('update teachers set short = ? where id = ?', (extra['short'], item_id))
            if table == 'subjects' and extra.get('short'):
                c.execute('update subjects set short = ? where id = ?', (extra['short'], item_id))
            if table == 'classes' and extra.get('label'):
                c.execute('update classes set label = ? where id = ?', (extra['label'], item_id))

        for item in items:
            if kind == 'lessons':
                class_id = resolve('classes', item['class'], {})
                subject_id = resolve('subjects', item['subject'], {})
                teacher_id = resolve('teachers', item['teacher'], {})
                row = c.execute(
                    'select id from lessons where class_id = ? and subject_id = ? and length = ?',
                    (class_id, subject_id, item['length']),
                ).fetchone()
                if row:
                    c.execute('update lessons set teacher_id = ?, count = ? where id = ?', (teacher_id, item['count'], row['id']))
                    updated += 1
                else:
                    c.execute(
                        'insert into lessons (class_id, subject_id, teacher_id, length, count) values (?, ?, ?, ?, ?)',
                        (class_id, subject_id, teacher_id, item['length'], item['count']),
                    )
                    added += 1
            else:
                resolve(kind, item['name'], item)
    return {'created': dict(created), 'linked': dict(linked), 'lessons_added': added, 'lessons_updated': updated}


def export_workbook(kind):
    wb = Workbook()
    ws = wb.active
    ws.title = kind
    if kind == 'lessons':
        ws.append(['Class', 'Teacher', 'Subject', 'Length', 'Lessons per week'])
        for r in query(
            'select c.name as c, t.name as t, s.name as s, l.length, l.count from lessons l '
            'join classes c on c.id = l.class_id join teachers t on t.id = l.teacher_id '
            'join subjects s on s.id = l.subject_id order by c.name, l.id'
        ):
            ws.append([r['c'], r['t'], r['s'], r['length'], r['count']])
    elif kind == 'teachers':
        ws.append(['Name', 'Short'])
        for r in query('select name, short from teachers order by id'):
            ws.append([r['name'], r['short']])
    elif kind == 'classes':
        ws.append(['Name', 'Label'])
        for r in query('select name, label from classes order by name'):
            ws.append([r['name'], r['label']])
    else:
        ws.append(['Name', 'Short'])
        for r in query('select name, short from subjects order by id'):
            ws.append([r['name'], r['short']])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
